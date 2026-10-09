"""F9 周报导出：复用report.py的docx组件生成统计周报（风险分布+质检指标+审查趋势）"""
from datetime import datetime, timedelta
from pathlib import Path

from docx import Document
from docx.shared import Pt

from . import db, qc
from .report import _add_page_watermark, _set_cell_text


def _recent_rows(days: int = 7):
    """近N天上传的合同行（风险分布与趋势取数）"""
    cutoff = (datetime.now() - timedelta(days=days)).strftime("%Y-%m-%d")
    rows = db.list_contract_rows()
    return [r for r in rows if (r.get("upload_time") or "")[:10] >= cutoff]


def _trend(rows):
    """近7天每日审查/上传数"""
    daily = {}
    for r in rows:
        d = (r.get("upload_time") or "")[:10]
        daily[d] = daily.get(d, 0) + 1
    out = []
    today = datetime.now().date()
    for i in range(6, -1, -1):
        d = (today - timedelta(days=i)).strftime("%Y-%m-%d")
        out.append({"date": d[5:], "count": daily.get(d, 0)})
    return out


def generate_weekly_report(out_path: Path) -> Path:
    """生成docx周报：①风险分布 ②质检指标（准确率/误报率）③审查趋势 ④审查人效率"""
    doc = Document()
    _add_page_watermark(doc)
    doc.add_heading("ClauseGuard 合同审查统计周报", level=0)
    p = doc.add_paragraph("统计周期：近7天　生成时间："
                          + datetime.now().strftime("%Y-%m-%d %H:%M"))
    p.runs[0].font.size = Pt(9)

    rows = _recent_rows()

    # 一、风险分布
    doc.add_heading("一、风险分布（近7天）", level=1)
    level_cnt, dim_cnt = {}, {}
    for r in rows:
        review = r.get("review")
        if not isinstance(review, dict):
            continue
        for f in review.get("findings", []):
            level_cnt[str(f.get("risk_level") or "未分级")] = \
                level_cnt.get(str(f.get("risk_level") or "未分级"), 0) + 1
            dim_cnt[str(f.get("dimension") or "其他")] = \
                dim_cnt.get(str(f.get("dimension") or "其他"), 0) + 1
    t = doc.add_table(rows=2, cols=max(len(level_cnt), 1) + 1)
    t.rows[0].cells[0].text = "风险等级"
    t.rows[1].cells[0].text = "发现数"
    for i, (k, v) in enumerate(sorted(level_cnt.items()), 1):
        t.rows[0].cells[i].text = k
        t.rows[1].cells[i].text = str(v)
    t2 = doc.add_table(rows=2, cols=max(len(dim_cnt), 1) + 1)
    t2.rows[0].cells[0].text = "维度"
    t2.rows[1].cells[0].text = "发现数"
    for i, (k, v) in enumerate(sorted(dim_cnt.items()), 1):
        t2.rows[0].cells[i].text = k
        t2.rows[1].cells[i].text = str(v)

    # 二、质检指标
    doc.add_heading("二、质检指标（各规则准确率/误报率）", level=1)
    stats = qc.rule_stats()
    if stats:
        t3 = doc.add_table(rows=1, cols=8)
        for i, h in enumerate(["规则", "维度", "抽样", "已复核", "正确", "错误",
                               "准确率%", "误报率%"]):
            _set_cell_text(t3.rows[0].cells[i], h, bold=True)
        for s in stats:
            row = t3.add_row()
            vals = [s["rule_name"][:30], s["dimension"] or "-", s["sampled_total"],
                    s["reviewed_total"], s["correct_total"], s["wrong_total"],
                    s["accuracy_rate"] if s["accuracy_rate"] is not None else "-",
                    s["false_positive_rate"] if s["false_positive_rate"] is not None else "-"]
            for i, v in enumerate(vals):
                _set_cell_text(row.cells[i], str(v))
    else:
        doc.add_paragraph("暂无质检抽样数据。")

    # 三、审查趋势
    doc.add_heading("三、审查趋势（近7天每日上传/审查量）", level=1)
    trend = _trend(rows)
    t4 = doc.add_table(rows=2, cols=8)
    t4.rows[0].cells[0].text = "日期"
    t4.rows[1].cells[0].text = "数量"
    for i, d in enumerate(trend, 1):
        t4.rows[0].cells[i].text = d["date"]
        t4.rows[1].cells[i].text = str(d["count"])

    # 四、审查人效率排行
    doc.add_heading("四、审查人效率排行", level=1)
    eff = qc.reviewer_efficiency()
    if eff:
        t5 = doc.add_table(rows=1, cols=3)
        for i, h in enumerate(["审查人", "处理量（次）", "平均耗时（秒）"]):
            _set_cell_text(t5.rows[0].cells[i], h, bold=True)
        for e in eff[:10]:
            row = t5.add_row()
            _set_cell_text(row.cells[0], str(e["username"]))
            _set_cell_text(row.cells[1], str(e["review_count"]))
            _set_cell_text(row.cells[2], str(e["avg_elapsed_s"] if e["avg_elapsed_s"] is not None else "-"))
    else:
        doc.add_paragraph("暂无审查人效率数据。")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
