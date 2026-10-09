"""F1功能自测：生成测试合同docx -> parser切分 -> reviewer审查 -> annotator批注回写 -> diff比对"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from docx import Document

BASE = Path(__file__).parent / "storage" / "f1_test"
BASE.mkdir(parents=True, exist_ok=True)


def make_contract(path: Path, version: int = 1):
    doc = Document()
    doc.add_heading("设备采购合同", level=1)
    doc.add_paragraph("甲方（买方）：北京示例科技有限公司")
    doc.add_paragraph("乙方（卖方）：上海供应设备有限公司")
    items = [
        "第一条 合同标的\n甲方向乙方采购服务器设备20台。",
        "第二条 合同金额\n合同总价为人民币50万元。",
        "第三条 交付期限\n乙方应于2026年12月31日前交付全部设备。",
        "第四条 质量标准\n设备应符合国家相关质量标准。",
        "第五条 违约责任\n任何一方违约应赔偿对方损失。",
        "第六条 争议解决\n本合同争议由合同签订地人民法院管辖。",
    ]
    if version == 2:
        items[1] = "第二条 合同金额\n合同总价为人民币80万元。"
        items[2] = "第三条 交付期限\n乙方应于2026年10月31日前交付全部设备。"
        items.append("第七条 保密条款\n双方对合同内容负有保密义务。")
    for it in items:
        head, _, body = it.partition("\n")
        doc.add_paragraph(head)
        doc.add_paragraph(body)
    doc.save(str(path))


def main():
    from app import parser, reviewer, annotator, diff
    from app.models import Contract, Clause, ExtractedInfo

    # 1.生成测试合同
    src1 = BASE / "合同v1.docx"
    src2 = BASE / "合同v2.docx"
    make_contract(src1, 1)
    make_contract(src2, 2)
    print("[1] 测试合同生成 OK:", src1.name, src2.name)

    # 2.parser切分
    text = parser.extract_text(src1)
    clauses = parser.split_clauses(text)
    print(f"[2] parser切分 OK: {len(clauses)}条，ids={[c['clause_id'] for c in clauses]}")

    # 3.reviewer审查
    extracted = None
    findings = reviewer.review(text, clauses, extracted)
    print(f"[3] reviewer审查 OK: {len(findings)}项发现")
    for f in findings[:5]:
        print("    -", f.get("risk_level"), f.get("clause_id"), f.get("rule_name"))

    # 4.annotator批注回写
    contract = Contract(
        contract_id="test01", filename="合同v1.docx", upload_time="2026-10-08",
        raw_text=text, clauses=[Clause(**c) for c in clauses],
        extracted=None,
    )
    if findings:
        from app.models import ReviewResult
        contract.review = ReviewResult(
            contract_id="test01", reviewed_at="2026-10-08 12:00:00",
            findings=findings, summary={"total": len(findings)})
    out_ann = annotator.generate_annotated(contract, BASE / "带批注版.docx", src1)
    out_clean = annotator.generate_clean(contract, BASE / "干净版.docx", src1)
    print("[4] 批注回写 OK:", out_ann.name, out_clean.name)

    # 验证批注已写入comments.xml
    import zipfile
    with zipfile.ZipFile(out_ann) as z:
        assert "word/comments.xml" in z.namelist(), "comments.xml缺失"
        cxml = z.read("word/comments.xml").decode("utf-8")
        n = cxml.count("<w:comment ")
        print(f"    comments.xml批注数: {n}")
        assert n >= 1, "批注未写入"
    with zipfile.ZipFile(out_clean) as z:
        assert "word/comments.xml" not in z.namelist(), "干净版不应含comments.xml"
    print("    干净版无批注 OK")

    # 5.diff比对
    text2 = parser.extract_text(src2)
    clauses2 = parser.split_clauses(text2)
    result = diff.diff_clauses(clauses, clauses2)
    print("[5] diff比对 OK:", result["summary"])
    for i in result["items"]:
        if i["change_type"] != "unchanged":
            print(f"    {i['change_type']} 第{i['clause_id']}条 高危={i['is_high_risk']} {i['new_title'] or i['old_title']}")

    print("\n=== 全部自测通过 ===")


if __name__ == "__main__":
    main()
